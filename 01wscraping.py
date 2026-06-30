import requests
from bs4 import BeautifulSoup
import pandas as pd
import re
from openpyxl import load_workbook

URL = "https://www.quiniela15.com/pronostico-quiniela"

headers = {
    "User-Agent": "Mozilla/5.0"
}

# Descargar página
html = requests.get(URL, headers=headers).text

# Parsear HTML
soup = BeautifulSoup(html, "html.parser")

texto = soup.get_text(" ", strip=True)

# Buscar porcentajes
patron = r'(Q15|LAE|APU):\s*(\d+)%\s*\|\s*(\d+)%\s*\|\s*(\d+)%'

coincidencias = re.findall(patron, texto)

partidos = []
grupo_actual = []

for fuente, uno, equis, dos in coincidencias:

    grupo_actual.append({
        "fuente": fuente,
        "1": int(uno),
        "X": int(equis),
        "2": int(dos)
    })

    # normalmente 3 fuentes por partido
    if len(grupo_actual) == 3:
        partidos.append(grupo_actual)
        grupo_actual = []

# si el último tiene 2 fuentes
if grupo_actual:
    partidos.append(grupo_actual)

# Crear filas
filas = []

for i, partido in enumerate(partidos, start=1):

    fila = {
        "Partido": i,
        "Q15_1": None,
        "Q15_X": None,
        "Q15_2": None,
        "LAE_1": None,
        "LAE_X": None,
        "LAE_2": None,
        "APU_1": None,
        "APU_X": None,
        "APU_2": None,
        "MEDIA_1": None,
        "MEDIA_X": None,
        "MEDIA_2": None
    }

    for dato in partido:

        fuente = dato["fuente"]

        fila[f"{fuente}_1"] = dato["1"]
        fila[f"{fuente}_X"] = dato["X"]
        fila[f"{fuente}_2"] = dato["2"]

    filas.append(fila)

# DataFrame
df = pd.DataFrame(filas)

nombre_excel = "quiniela_pronosticos.xlsx"

# Guardar Excel
df.to_excel(nombre_excel, index=False)

# Abrir workbook para insertar fórmulas
wb = load_workbook(nombre_excel)
ws = wb.active

# Columnas Excel:
# B = Q15_1
# E = LAE_1
# H = APU_1
#
# C = Q15_X
# F = LAE_X
# I = APU_X
#
# D = Q15_2
# G = LAE_2
# J = APU_2

for fila in range(2, ws.max_row + 1):

    ws[f"K{fila}"] = f'=PROMEDIO(B{fila},E{fila},H{fila})'
    ws[f"L{fila}"] = f'=PROMEDIO(C{fila},F{fila},I{fila})'
    ws[f"M{fila}"] = f'=PROMEDIO(D{fila},G{fila},J{fila})'

# Guardar
wb.save(nombre_excel)

print(f"Excel generado: {nombre_excel}")